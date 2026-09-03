"""Generate deterministic PDF, XLSX, and scan-style PNG quotation samples."""

from __future__ import annotations

import json
import random
import re
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas


OUTPUT_DIR = Path(__file__).resolve().parent
FIXED_TIME = datetime(2026, 9, 2, tzinfo=timezone.utc)

QUOTE_FIXTURES = {
    "vendor_a_quote.pdf": {
        "vendor": "Apex Bearing Works",
        "item": "Bearing 6205",
        "quantity": 500,
        "unit_price_inr": 412,
        "freight_inr": 0,
        "tax_percent": 18,
        "total_landed_cost_inr": 206000,
        "lead_time_days": 12,
        "payment_terms": "Net 30",
        "warranty_months": 12,
    },
    "vendor_b_quote.xlsx": {
        "vendor": "Beacon Industrial Components",
        "item": "Bearing 6205",
        "quantity": 500,
        "unit_price_inr": 438,
        "freight_inr": 0,
        "tax_percent": 18,
        "total_landed_cost_inr": 219000,
        "lead_time_days": 7,
        "payment_terms": "Net 15",
        "warranty_months": None,
    },
    "vendor_c_quote.png": {
        "vendor": "Crest Motion Supplies",
        "item": "Bearing 6205",
        "quantity": 500,
        "unit_price_inr": 429,
        "freight_inr": 0,
        "tax_percent": 18,
        "printed_subtotal_inr": 211500,
        "calculated_subtotal_inr": 214500,
        "total_landed_cost_inr": 214500,
        "lead_time_days": 10,
        "payment_terms": "Net 45",
        "warranty_months": 18,
    },
}

EXPECTED_FINDINGS = {
    "recommended_vendor": "Apex Bearing Works",
    "review_items": [
        "Vendor B warranty is not stated.",
        "Vendor C printed subtotal INR 211,500 conflicts with 500 x INR 429 = INR 214,500.",
    ],
    "notes": "Tax is recoverable and excluded from the landed-cost comparison in all samples.",
}


def _generate_pdf(path: Path) -> None:
    quote = QUOTE_FIXTURES[path.name]
    pdf = canvas.Canvas(str(path), pagesize=A4, pageCompression=1, invariant=1)
    width, height = A4
    pdf.setTitle("Vendor A quotation")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(24 * mm, height - 25 * mm, quote["vendor"])
    pdf.setFont("Helvetica", 10)
    lines = [
        "QUOTATION A-6205-0902",
        "Date: 02 September 2026",
        "Customer: Meridian Industrial Supply",
        "",
        f"Item: {quote['item']}",
        f"Quantity: {quote['quantity']} units",
        f"Unit price: INR {quote['unit_price_inr']}",
        f"Freight: Included (INR {quote['freight_inr']})",
        f"Subtotal / landed cost before recoverable tax: INR {quote['total_landed_cost_inr']:,}",
        f"GST: {quote['tax_percent']}% extra; recoverable input tax, excluded from comparison",
        f"Lead time: {quote['lead_time_days']} calendar days",
        f"Payment terms: {quote['payment_terms']}",
        f"Warranty: {quote['warranty_months']} months from delivery",
        "Validity: 30 days",
    ]
    y = height - 42 * mm
    for line in lines:
        pdf.drawString(24 * mm, y, line)
        y -= 7 * mm
    pdf.setFont("Helvetica-Oblique", 8)
    pdf.drawString(24 * mm, 18 * mm, "Synthetic Phase 0 fixture - no real vendor or transaction")
    pdf.showPage()
    pdf.save()


def _generate_xlsx(path: Path) -> None:
    quote = QUOTE_FIXTURES[path.name]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quotation"
    workbook.properties.creator = "BlaiseLogic Project009"
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)

    rows = [
        ("VENDOR QUOTATION", quote["vendor"]),
        ("Quotation number", "BIC-6205-26"),
        ("Date", "02 September 2026"),
        ("Customer", "Meridian Industrial Supply"),
        ("Item", quote["item"]),
        ("Quantity", quote["quantity"]),
        ("Unit price (INR)", quote["unit_price_inr"]),
        ("Freight (INR)", quote["freight_inr"]),
        ("Landed cost before recoverable tax (INR)", quote["total_landed_cost_inr"]),
        ("GST (%)", quote["tax_percent"]),
        ("Lead time (calendar days)", quote["lead_time_days"]),
        ("Payment terms", quote["payment_terms"]),
        ("Warranty", ""),
        ("Validity", "21 days"),
        ("Note", "Warranty was not supplied; confirm before award."),
    ]
    for row_index, (label, value) in enumerate(rows, 1):
        sheet.cell(row=row_index, column=1, value=label)
        sheet.cell(row=row_index, column=2, value=value)
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["B1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["B1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.column_dimensions["A"].width = 43
    sheet.column_dimensions["B"].width = 38
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)
    _normalize_xlsx_archive(path)


def _normalize_xlsx_archive(path: Path) -> None:
    """Remove ZIP member timestamps so identical fixture data yields identical bytes."""

    normalized_path = path.with_suffix(".normalized.xlsx")
    with ZipFile(path, "r") as source, ZipFile(
        normalized_path, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for source_info in sorted(source.infolist(), key=lambda item: item.filename):
            payload = source.read(source_info.filename)
            if source_info.filename == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    rb"\g<1>2026-09-02T00:00:00Z\g<2>",
                    payload,
                )
            target_info = ZipInfo(
                filename=source_info.filename,
                date_time=(2026, 9, 2, 0, 0, 0),
            )
            target_info.compress_type = source_info.compress_type
            target_info.comment = source_info.comment
            target_info.external_attr = source_info.external_attr
            target_info.internal_attr = source_info.internal_attr
            target_info.create_system = source_info.create_system
            destination.writestr(target_info, payload)
    normalized_path.replace(path)


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for name in ("DejaVuSansMono.ttf", "arial.ttf"):
        try:
            return ImageFont.truetype(name, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _generate_png(path: Path) -> None:
    quote = QUOTE_FIXTURES[path.name]
    image = Image.new("RGB", (1500, 1900), "#f8f5ec")
    draw = ImageDraw.Draw(image)
    title_font = _font(38)
    body_font = _font(25)
    small_font = _font(20)
    lines = [
        quote["vendor"].upper(),
        "SCANNED QUOTATION C-0902",
        "Date: 02 September 2026",
        "Customer: Meridian Industrial Supply",
        "",
        f"Item                       {quote['item']}",
        f"Quantity                   {quote['quantity']} units",
        f"Unit price                 INR {quote['unit_price_inr']}",
        f"Printed subtotal           INR {quote['printed_subtotal_inr']:,}",
        f"Total landed cost          INR {quote['total_landed_cost_inr']:,}",
        f"Freight                    Included",
        f"GST                        {quote['tax_percent']}% extra / recoverable",
        f"Lead time                  {quote['lead_time_days']} calendar days",
        f"Payment terms              {quote['payment_terms']}",
        f"Warranty                   {quote['warranty_months']} months",
        "",
        "Please verify totals against quantity x unit price.",
    ]
    y = 130
    for index, line in enumerate(lines):
        font = title_font if index == 0 else body_font
        draw.text((130, y), line, font=font, fill="#252525")
        y += 74 if index == 0 else 64
    draw.rectangle((90, 90, 1410, 1680), outline="#77736c", width=3)
    draw.text(
        (130, 1740),
        "Synthetic Phase 0 fixture - no real vendor or transaction",
        font=small_font,
        fill="#565656",
    )

    random.seed(9009)
    for _ in range(550):
        x = random.randrange(image.width)
        y = random.randrange(image.height)
        shade = random.randrange(205, 240)
        draw.point((x, y), fill=(shade, shade, shade))
    image = image.rotate(0.65, resample=Image.Resampling.BICUBIC, expand=False, fillcolor="white")
    image.save(path, format="PNG", optimize=True)


def generate_samples(output_dir: Path = OUTPUT_DIR) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = [
        output_dir / "vendor_a_quote.pdf",
        output_dir / "vendor_b_quote.xlsx",
        output_dir / "vendor_c_quote.png",
    ]
    _generate_pdf(outputs[0])
    _generate_xlsx(outputs[1])
    _generate_png(outputs[2])
    (output_dir / "expected_results.json").write_text(
        json.dumps(
            {"quotes": QUOTE_FIXTURES, "expected_findings": EXPECTED_FINDINGS},
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return outputs


if __name__ == "__main__":
    for generated in generate_samples():
        print(generated)
